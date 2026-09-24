import json, re, sys, collections, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
S=sys.argv[1]; LV=[int(x) for x in sys.argv[2].split(',')]
tb=json.load(open(S+'/tbcl.json',encoding='utf-8'))
CAT={1:'個人資料',2:'日常起居',3:'職業',4:'休閒、娛樂',5:'交通、旅遊',6:'社交、人際',7:'身體、醫療',8:'教育、學習',9:'購物、商店',10:'餐飲、烹飪',11:'公共服務',12:'安全',13:'自然環境',14:'社會',15:'文化',16:'情緒、態度',17:'科技'}
TSUB={'數':'angka/jumlah','時':'waktu','代':'kata ganti/penunjuk','量':'kata bantu bilangan','問':'kata tanya','方':'arah/posisi','呼':'sapaan','虛':'kata fungsi/adverbia/modal'}
# ledger: old module per word
led={}
for line in open(r'C:\Users\user\OneDrive\文件\TOCFL\ledger.md',encoding='utf-8'):
    m=re.match(r'\|\s*([A-Z]\d+)\s*\|[^|]*\|[^|]*\|([^|]*)\|([^|]*)\|',line)
    if m:
        for w in re.split(r'[,、，]\s*',m.group(2)): led.setdefault(w.strip(),m.group(1)+' 核心')
        for w in re.split(r'[,、，]\s*',m.group(3)): led.setdefault(w.strip(),m.group(1)+' 補充')
def oldmod(w):
    for part in [w]+w.split('/'):
        if part in led: return led[part]
    return ''
wb=openpyxl.Workbook(); wb.remove(wb.active)
summ=wb.create_sheet('Ringkasan')
hdrfill=PatternFill('solid',fgColor='1F4E78'); flagfill=PatternFill('solid',fgColor='FFF2CC'); offfill=PatternFill('solid',fgColor='E2EFDA')
summ_rows=[]
for lv in LV:
    mp={}
    for line in open(f'{S}/map_L{lv}.txt',encoding='utf-8'):
        if line.startswith('#') or not line.strip(): continue
        p=line.rstrip('\n').split(' ',2); mp[int(p[0])]=(p[1],p[2] if len(p)>2 else '')
    words=[o for o in tb if o['lv']==lv]
    miss=[o['id'] for o in words if o['cat']=='核心詞' and o['id'] not in mp]
    extra=[i for i in mp if i not in {o['id'] for o in words if o['cat']=='核心詞'}]
    print('L',lv,'missing',miss,'extra',extra)
    ws=wb.create_sheet(f'Level {lv}')
    H=['序號','詞語','拼音','級別','情境 resmi (總表)','Usulan Claude','Nama kategori / sub-pool','Catatan & 義項','⚑ Perlu keputusan','Modul lama (ledger)','Keputusan Carli (OK / ganti ke…)']
    ws.append(H)
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=hdrfill; c.alignment=Alignment(wrap_text=True,vertical='center')
    cnt=collections.Counter()
    for o in words:
        if o['cat']!='核心詞':
            n=int(o['cat'].split('.')[0]); code=str(n); name=CAT[n]; note='(kategori resmi, tidak diubah)'; flag=''; src='resmi'
        else:
            code,note=mp[o['id']]; flag='⚑' if code.endswith('!') else ''; code=code.rstrip('!')
            name=('工具詞 · '+TSUB[code[1]]) if code.startswith('T') else CAT[int(code)]; src='usulan'
        cnt[code if code.startswith('T') else int(code)]+=1
        ws.append([o['id'],o['w'],o['py'],f"第{lv}{'*' if o['star'] else ''}級",o['cat'],code,name,note,flag,oldmod(o['w']),''])
        r=ws.max_row
        if src=='resmi':
            for c in ws[r]: c.fill=offfill
        elif flag:
            for c in ws[r]: c.fill=flagfill
    for col,wd in zip('ABCDEFGHIJK',[7,16,16,8,16,10,24,40,10,14,26]): ws.column_dimensions[col].width=wd
    ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
    summ_rows.append((lv,cnt,len(words)))
summ.append(['Kategori']+[f'Level {lv}' for lv,_,_ in summ_rows])
for c in summ[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=hdrfill
keys=list(range(1,18))+['T'+k for k in TSUB]
for k in keys:
    nm=f'{k}. {CAT[k]}' if isinstance(k,int) else f'{k} 工具詞 · {TSUB[k[1]]}'
    summ.append([nm]+[cnt.get(k,0) for _,cnt,_ in summ_rows])
summ.append(['TOTAL']+[t for _,_,t in summ_rows])
summ.column_dimensions['A'].width=40
summ.append([]); summ.append(['Hijau = kategori resmi 總表 (tidak diubah). Kuning ⚑ = 多義詞/ambigu, mohon Carli putuskan. Putih = usulan Claude untuk kata berlabel 核心詞.'])
out=r'C:\Users\user\OneDrive\文件\TOCFL\pemetaan-kategori.xlsx'; wb.save(out)
for lv,cnt,t in summ_rows:
    print('Level',lv,'total',t, dict(sorted(((str(k),v) for k,v in cnt.items()))))
