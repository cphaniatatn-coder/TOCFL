"""Edit soal tes awal/akhir modul mini lewat Excel.

    py _kerja/mini/tes_excel.py ekspor   → buat _kerja/mini/edit-tes-mini.xlsx dari tes_A0/A1/A2.json
    py _kerja/mini/tes_excel.py impor    → tulis perubahan Excel kembali ke tes_*.json, lalu periksa kosakata,
                                           rekam audio baru, dan rakit ulang repo mini

Sumber tetap tes_A{0,1,2}.json; Excel hanya "jendela edit" (file yang isinya tidak berubah tidak ditulis ulang).
Bunyi bel 聽力: baris 🔔 di kolom "naskah audio" → disimpan sebagai t["bel"] (dibaca TocflAudio di web/js/studi.js).
"""
import json, os, re, subprocess, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.stdout.reconfigure(encoding='utf-8')
M = os.path.dirname(os.path.abspath(__file__)); K = os.path.dirname(M); R = os.path.dirname(K)
sys.path.insert(0, K)
XLSX = f'{M}/edit-tes-mini.xlsx'
LEVEL = ['A0', 'A1', 'A2']
ABC = 'ABCDEF'
BEL = '🔔'
OPSI = '（A）（B）（C）'   # penanda di naskah Part 1: di sinilah pilihan A, B, C dibacakan
TIPE = {'listen_pic': '聽力 Part 1', 'listen_dialog': '聽力 Part 2–4', 'read_sent': '閱讀 Part 1', 'read_pick': '閱讀 Part 2',
        'read_gap': '閱讀 Part 3', 'cloze': '閱讀 Part 4', 'read_mc': '閱讀 Part 5'}
KOLOM = ['level', 'paket', 'no', 'tipe', 'bagian', 'tema', 'naskah audio (聽力; 🔔 = bel)', 'gambar soal (ikon | keterangan)',
         'teks bacaan (閱讀)', '問 (閱讀 Part 5)', 'pilihan (satu per baris)', 'jawaban', 'alasan']
LEBAR = [7, 7, 5, 14, 13, 24, 46, 34, 44, 26, 40, 10, 44]
WARNA = {'A0': 'E8F5E9', 'A1': 'E3F2FD', 'A2': 'FFF8E1'}


def baca(lv): return json.load(open(f'{M}/tes_{lv}.json', encoding='utf-8'))

def tulis(lv, obj):
    s = json.dumps(obj, ensure_ascii=False, indent=2).replace('\n', '\r\n') + '\r\n'
    open(f'{M}/tes_{lv}.json', 'w', encoding='utf-8', newline='').write(s)

def bel_bawaan(t): return [len(t.get('lines', []))] if t['type'] == 'listen_dialog' else []

def sel(v): return '' if v is None else str(v).strip()
def daftar(s): return [x.strip() for x in sel(s).split('\n') if x.strip()]
def opsi_teks(o): return f"{o['icon']} | {o['label']}" if isinstance(o, dict) else o


# ---------- ekspor ----------
def naskah(t):
    if t['type'] == 'listen_pic':
        klip = [f"問：{t['question']}", OPSI]
        pos = {0: 0, 1: 1, 4: 2}           # nomor klip app → posisi baris (A/B/C dihitung satu blok)
    elif t['type'] == 'listen_dialog':
        klip = [f"{l['sp']}：{l['zh']}" for l in t['lines']] + [f"問：{t['question']}"]
        pos = {i: i for i in range(len(klip) + 1)}
    else:
        return ''
    out = []
    bel = t.get('bel', bel_bawaan(t))
    for i, k in enumerate(klip + [None]):
        if any(pos.get(b) == i for b in bel): out.append(BEL)
        if k: out.append(k)
    return '\n'.join(out)

def ekspor():
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'Petunjuk'
    for l in PETUNJUK.strip('\n').split('\n'): ws.append([l])
    ws.column_dimensions['A'].width = 120
    ws['A1'].font = Font(bold=True, size=14)
    ws = wb.create_sheet('Soal'); ws.append(KOLOM)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDDDDD')
    for i, w in enumerate(LEBAR): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    n = 0
    for lv in LEVEL:
        for paket, items in baca(lv).items():
            for t in items:
                pil = t['options']
                jaw = ', '.join(ABC[a] for a in t['answers']) if t['type'] == 'cloze' else ABC[t['answer']]
                ws.append([lv, paket, t['no'], t['type'], t['part'], t.get('tema', ''), naskah(t),
                           opsi_teks(t['picture']) if t.get('picture') else '', t.get('text', ''),
                           t.get('question', '') if t['type'] == 'read_mc' else '',
                           '\n'.join(opsi_teks(o) for o in pil), jaw, t.get('why', '')])
                fill = PatternFill('solid', fgColor=WARNA[lv])
                for c in ws[ws.max_row]: c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical='top')
                n += 1
    ws.freeze_panes = 'D2'; ws.auto_filter.ref = ws.dimensions
    wb.save(XLSX)
    print(f'Selesai: {XLSX}\n  Soal: {n} baris (A0/A1/A2 × paket A/B × 20)')


# ---------- impor ----------
def impor():
    import buat_audio as ba
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['Soal']
    head = [sel(c.value) for c in ws[1]]
    if head != KOLOM: sys.exit('✗ Judul kolom sheet "Soal" berubah. Jangan ubah baris pertama.')
    salah, catatan = [], []
    baru = {lv: {} for lv in LEVEL}
    lama = {lv: baca(lv) for lv in LEVEL}
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        v = dict(zip(KOLOM, (sel(x) for x in row)))
        if not any(v.values()): continue
        tempat = f'baris {r}'
        def err(m): salah.append(f'{tempat}: {m}')
        lv, paket, tipe = v['level'], v['paket'], v['tipe']
        if lv not in LEVEL: err(f'level "{lv}" harus A0/A1/A2'); continue
        if paket not in ('A', 'B'): err(f'paket "{paket}" harus A atau B'); continue
        try: no = int(float(v['no']))
        except ValueError: err('kolom no harus angka'); continue
        tempat = f'baris {r} ({lv}{paket}-{no})'
        if tipe not in TIPE: err(f'tipe "{tipe}" tidak dikenal ({", ".join(TIPE)})'); continue
        t = {'no': no, 'type': tipe, 'part': v['bagian'] or TIPE[tipe], 'tema': v['tema']}
        if v['gambar soal (ikon | keterangan)']:
            g = v['gambar soal (ikon | keterangan)'].split(' | ', 1)
            if len(g) != 2: err('gambar soal harus berbentuk  ikon | keterangan')
            else: t['picture'] = {'label': g[1].strip(), 'icon': g[0].strip()}
        if tipe.startswith('listen'):
            klip, bel, lines, q = 0, [], [], None
            for ln in daftar(v['naskah audio (聽力; 🔔 = bel)']):
                if ln.replace('️', '') == BEL.replace('️', '') or ln.startswith(BEL):
                    bel.append(klip); continue
                if ln.startswith('問：') or ln.startswith('問:'):
                    q = ln[2:].strip(); klip += 1; continue
                if tipe == 'listen_pic' and re.fullmatch(r'[（(]?A[）)]?\s*[（(]?B[）)]?\s*[（(]?C[）)]?', ln.replace(' ', '')):
                    klip += 3; continue
                m = re.match(r'([^：:]{1,8})[：:](.+)', ln)
                if tipe == 'listen_dialog' and m:
                    sp = m.group(1).strip()
                    if sp not in ba.PEMBICARA: catatan.append(f'{tempat}: pembicara "{sp}" belum ada di SPEAKERS (js/media.js) → suara F1')
                    lines.append({'sp': sp, 'zh': m.group(2).strip()}); klip += 1; continue
                err(f'baris naskah tidak dikenali: "{ln}" (format  女：kalimat  /  問：pertanyaan  /  🔔)')
            if not q: err('naskah audio harus punya baris  問：pertanyaan')
            if tipe == 'listen_dialog':
                if not lines: err('naskah dialog kosong')
                t['lines'] = lines
            t['question'] = q
            # nomor bel = jumlah klip sebelum 🔔 (Part 1: 問=0, A/B/C=1–3, akhir=4); simpan hanya bila beda dari bawaan
            bel = sorted(set(bel))
            if bel != bel_bawaan(t): t['bel'] = bel
        else:
            if tipe != 'read_pick' and not v['teks bacaan (閱讀)']: err('teks bacaan kosong')
            if v['teks bacaan (閱讀)']: t['text'] = v['teks bacaan (閱讀)']
            if tipe == 'read_mc':
                if not v['問 (閱讀 Part 5)']: err('kolom 問 kosong')
                t['question'] = v['問 (閱讀 Part 5)']
        pil = daftar(v['pilihan (satu per baris)'])
        t['options'] = [{'label': o.split(' | ', 1)[1].strip(), 'icon': o.split(' | ', 1)[0].strip()} if ' | ' in o else o for o in pil]
        if len(pil) < 3: err('pilihan minimal 3')
        huruf = [x.strip().upper() for x in re.split(r'[,，、\s]+', v['jawaban']) if x.strip()]
        if any(len(h) != 1 or h not in ABC[:len(pil)] for h in huruf) or not huruf:
            err(f'jawaban "{v["jawaban"]}" harus huruf pilihan (A–{ABC[max(len(pil), 1) - 1]})')
        elif tipe == 'cloze':
            t['answers'] = [ABC.index(h) for h in huruf]
            n_kosong = len(re.findall(r'（\d）', t.get('text', '')))
            if n_kosong != len(huruf): err(f'teks punya {n_kosong} titik kosong （1）… tetapi jawaban {len(huruf)} huruf')
        else:
            if len(huruf) != 1: err('jawaban harus satu huruf')
            t['answer'] = ABC.index(huruf[0])
        if tipe == 'read_gap' and '（　）' not in t.get('text', '') and '（ ）' not in t.get('text', ''):
            err('teks Part 3 harus punya titik kosong （　）')
        t['why'] = v['alasan']
        if no in baru[lv].setdefault(paket, {}): err('nomor soal dobel')
        baru[lv][paket][no] = t

    if salah:
        print('✗ Ada kesalahan — tidak ada file yang diubah:'); print('\n'.join('  ' + s for s in salah)); sys.exit(1)

    berubah = []
    for lv in LEVEL:
        obj = {}
        for paket in ('A', 'B'):
            items = []
            for no in sorted(baru[lv].get(paket, {})):
                t = baru[lv][paket][no]
                old = next((x for x in lama[lv].get(paket, []) if x['no'] == no), None)
                if old:   # urutan kunci mengikuti file lama supaya diff kecil
                    t = {**{k: t[k] for k in old if k in t}, **t}
                    if t['type'] == 'listen_pic' and old.get('picture') != t.get('picture'):
                        catatan.append(f'{lv}{paket}-{no}: gambar Part 1 berubah → minta Claude menggambar ulang (gambar_tes.py)')
                items.append(t)
            obj[paket] = items
        if obj != lama[lv]:
            tulis(lv, obj); berubah.append(lv)
    if not berubah:
        print('Tidak ada perubahan.'); return
    print('✓ Ditulis: ' + ', '.join(f'tes_{lv}.json' for lv in berubah))
    for c in catatan: print('  ⚑ ' + c)

    py = sys.executable
    def jalan(cmd, judul):
        print(f'\n== {judul}')
        r = subprocess.run([py] + cmd, cwd=R, capture_output=True, text=True, encoding='utf-8', env=dict(os.environ, PYTHONIOENCODING='utf-8'))
        print((r.stdout + r.stderr).strip() or '(tidak ada masalah)')
        return r.returncode == 0
    jalan([f'{M}/cek_tes.py'], 'Cek kosakata & grammar sesuai level')
    jalan([f'{K}/buat_audio.py'], 'Audio (kalimat baru direkam; butuh internet)')
    jalan([f'{K}/buat_mini.py'], 'Rakit repo mini (teks baru yang belum diterjemahkan → _kerja/mini/kurang_*.txt)')
    print('\n✓ Selesai. Buka tocfl-mini untuk mengecek, lalu commit & push.')


PETUNJUK = """
EDIT SOAL TES AWAL & AKHIR (MODUL MINI) LEWAT EXCEL

1. Buat file ini (selalu dari data terbaru):   py _kerja/mini/tes_excel.py ekspor
2. Edit sheet "Soal", SIMPAN, lalu TUTUP file-nya.
3. Masukkan perubahan:                         py _kerja/mini/tes_excel.py impor
   (memeriksa isi, merekam audio baru, merakit ulang tocfl-mini). Kalau ada kesalahan, tidak ada file yang diubah.

Paket A = tes awal (pre-test), paket B = tes akhir (post-test). Warna baris: hijau A0, biru A1, kuning A2.
Jangan ubah judul kolom (baris 1). Untuk membuat baris baru di dalam satu sel: Alt+Enter.

NASKAH AUDIO 聽力 — satu baris per bunyi, dibacakan dari atas ke bawah:
  女：這件衣服多少錢？     ← baris dialog (pembicara：kalimat, titik dua lebar ：)
  男：六百五十塊。
  🔔                       ← BUNYI BEL di sini (salin ikon ini; boleh lebih dari satu)
  問：這件衣服多少錢？     ← pertanyaan yang dibacakan narator
Part 1 (gambar):  問：…  lalu  （A）（B）（C）  = tempat pilihan A, B, C dibacakan. 🔔 boleh sebelum 問, sebelum
  （A）（B）（C）, atau di akhir.
Tanpa 🔔 sama sekali = tidak ada bel. Teks 問 tidak tampil di layar selama tes (hanya dibacakan).

KOLOM LAIN
- gambar soal / pilihan bergambar:  ikon | keterangan   (mis.  🏷️ | label harga NT$650)
- pilihan: satu per baris. jawaban: huruf (A, B, C, D). Part 4 (cloze): huruf per titik kosong, urut, mis.  A, B, C, D, E
- Part 3: teks harus punya titik kosong （　）. Part 4: titik kosong ditulis （1）（2）…
- Kalimat Mandarin hanya boleh memakai kosakata sampai level itu (A0 ≤ A25, A1 ≤ B46, A2 ≤ C67) — dicek otomatis.
- Gambar Part 1 adalah ilustrasi buatan kode: kalau isi gambarnya berubah, minta Claude menggambar ulang.
- Keterangan gambar & alasan dalam bahasa Indonesia; terjemahan English/Tiếng Việt untuk teks baru dibuat terpisah.
"""

if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else ''
    if cmd == 'ekspor': ekspor()
    elif cmd == 'impor': impor()
    else: print(__doc__)
