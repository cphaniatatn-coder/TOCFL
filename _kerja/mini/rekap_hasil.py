"""Rekap hasil uji coba modul mini → Excel yang rapi untuk dibaca & dianalisis.

Pakai:  py rekap_hasil.py  [path/測試結果.xlsx]
  - Input default: 測試結果.xlsx di folder tocfl-mini (unduh dari Google Sheet: File → Download → .xlsx).
  - Output: Rekap-測試結果.xlsx di folder yang sama (berisi email peserta → JANGAN di-commit).
Kolom "Pakai?" di tab Ringkasan diisi otomatis (Ya/Tidak) menurut aturan di bawah, tetapi boleh diubah
manual di Excel; semua rata-rata ikut berubah.
"""
import json, sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

MINI = Path(__file__).resolve().parents[3] / 'tocfl-mini'
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else MINI / '測試結果.xlsx'
OUT = SRC.with_name('Rekap-' + SRC.stem + '.xlsx')
# Commit f4c9c60: sebelum ini kunci tes kebanyakan A dan urutan pilihan tidak diacak.
TES_BARU = datetime(2026, 10, 8, 0, 19)

KUES = json.load(open(MINI / 'data/id/studi.json', encoding='utf-8'))['kuesioner']
BUTIR = [(d['dim'], it['t'], bool(it.get('r'))) for d in KUES['cemas'] for it in d['items']]
EVAL = KUES['eval']

# ---------- baca data ----------
wb_in = openpyxl.load_workbook(SRC, data_only=True)
def tab(nama):
    if nama not in wb_in.sheetnames: return []
    rows = [r for r in wb_in[nama].iter_rows(values_only=True) if any(v is not None for v in r)]
    if not rows: return []
    h = [str(x) for x in rows[0]]
    return [dict(zip(h, r)) for r in rows[1:]]
def tgl(v):
    if isinstance(v, datetime): return v
    if not v: return None
    s = str(v).replace('上午', 'AM').replace('下午', 'PM')
    for f in ('%Y/%m/%d %p %I:%M:%S', '%Y/%m/%d', '%Y-%m-%d %H:%M:%S'):
        try: return datetime.strptime(s, f)
        except ValueError: pass
    return None
def angka(v):
    if isinstance(v, datetime): return (v - datetime(1899, 12, 30)).days   # Sheets kadang memformat angka jadi tanggal
    try: return float(v) if v not in (None, '') else None
    except ValueError: return None

peserta, cemas, tes, modul, evals = tab('Peserta'), tab('Kecemasan'), tab('Tes'), tab('Modul'), tab('Evaluasi')
cemas_by = {(r['uid'], r['fase']): r for r in cemas}
tes_by = {(r['uid'], r['fase']): r for r in tes}
eval_by = {r['uid']: r for r in evals}
seragam = lambda vals: len([v for v in vals if v is not None]) > 1 and len(set(v for v in vals if v is not None)) == 1

# ---------- aturan "Pakai?" ----------
pertama = {}   # email → kode_asal yang paling awal mulai
for p in sorted(peserta, key=lambda p: tgl(p['mulai']) or datetime.max):
    pertama.setdefault((p['email'] or '').strip().lower(), p['kode_asal'])
kode_per_email = defaultdict(set)
for p in peserta: kode_per_email[(p['email'] or '').strip().lower()].add(p['kode_asal'])

def catatan(p):
    uid, alasan, info = p['uid'], [], []
    em = (p['email'] or '').strip().lower()
    if int(angka(p['putaran']) or 1) > 1: alasan.append('putaran 2 (bukan analisis utama)')
    if not (eval_by.get(uid) and tes_by.get((uid, 'post'))): alasan.append('belum selesai')
    if len(kode_per_email[em]) > 1:
        if pertama[em] != p['kode_asal']: alasan.append('email sama dengan baris lain (perangkat ke-2)')
        else: info.append('email ini juga dipakai di perangkat lain')
    sr = [n for n, vals in (('kuesioner awal', [cemas_by.get((uid, 'pre'), {}).get(f'b{i}') for i in range(1, 21)]),
                            ('kuesioner akhir', [cemas_by.get((uid, 'post'), {}).get(f'b{i}') for i in range(1, 21)]),
                            ('evaluasi', [eval_by.get(uid, {}).get(f'e{i}') for i in range(1, 11)])) if seragam(vals)]
    if sr: alasan.append('jawaban seragam: ' + ', '.join(sr))
    if any((tgl(tes_by.get((uid, f), {}).get('selesai')) or TES_BARU) < TES_BARU for f in ('pre', 'post')):
        alasan.append('tes versi lama (kunci kebanyakan A)')
    return '; '.join(alasan + info), 'Tidak' if alasan else 'Ya'

def status(p):
    uid = p['uid']
    if eval_by.get(uid): return 'Selesai'
    if tes_by.get((uid, 'post')): return 'Belum isi evaluasi'
    if not tes_by.get((uid, 'pre')): return 'Belum tes awal'
    return f"Belajar · {p['posisi_terakhir'] or '-'}"

# ---------- gaya ----------
F = 'Arial'
HEAD = PatternFill('solid', fgColor='1F4E78'); HFONT = Font(name=F, bold=True, color='FFFFFF')
INPUT = PatternFill('solid', fgColor='FFF2CC'); ABU = PatternFill('solid', fgColor='F2F2F2')
HIJAU = PatternFill('solid', fgColor='C6EFCE'); MERAH = PatternFill('solid', fgColor='FFC7CE')
tipis = Side(style='thin', color='BFBFBF'); KOTAK = Border(top=tipis, bottom=tipis, left=tipis, right=tipis)

wb = openpyxl.Workbook()
def sheet(nama, kolom, baris, lebar=None, mulai=1):
    ws = wb.create_sheet(nama)
    for j, k in enumerate(kolom, 1):
        c = ws.cell(mulai, j, k[0] if isinstance(k, tuple) else k)
        c.fill, c.font, c.border = HEAD, HFONT, KOTAK
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        if isinstance(k, tuple) and k[1]: c.comment = Comment(k[1], 'rekap', width=320, height=110)
    for i, r in enumerate(baris, mulai + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v); c.font = Font(name=F); c.border = KOTAK
    ws.freeze_panes = ws.cell(mulai + 1, 3)
    ws.auto_filter.ref = f'A{mulai}:{L(len(kolom))}{max(mulai + 1, mulai + len(baris))}'
    ws.row_dimensions[mulai].height = 32
    for j in range(1, len(kolom) + 1): ws.column_dimensions[L(j)].width = (lebar or {}).get(j, 11)
    return ws
def fmt(ws, kol, f, r0, r1):
    for r in range(r0, r1 + 1): ws[f'{kol}{r}'].number_format = f

# ---------- Kecemasan ----------
kc = sorted(cemas, key=lambda r: (r['uid'], r['fase'] != 'pre'))
k_kol = ['Kunci', 'Nama', 'uid', 'Putaran', 'Level', 'Fase', 'Waktu', ('Total', '20–100. Makin tinggi = makin cemas. b10 sudah dibalik.'),
         ('A umum', BUTIR[0][0]), ('B dengar', BUTIR[5][0]), ('C baca', BUTIR[10][0]), ('D kesiapan', BUTIR[15][0])] + \
        [(f'b{i}', f"{d}\n\n{t}" + ('\n\n(R) butir terbalik: nilai di sini MENTAH, di Total sudah dibalik (6 − nilai).' if r else ''))
         for i, (d, t, r) in enumerate(BUTIR, 1)] + ['Seragam?', 'Pakai?']
n = len(kc) + 1
k_rows = [[f"{r['uid']}|{r['fase']}", r['nama'], r['uid'], angka(r['putaran']), r['level'], 'awal' if r['fase'] == 'pre' else 'akhir',
           tgl(r['waktu']), angka(r['total']), angka(r['A_umum']), angka(r['B_dengar']), angka(r['C_baca']), angka(r['D_kesiapan'])] +
          [angka(r[f'b{i}']) for i in range(1, 21)] +
          ['Ya' if seragam([r[f'b{i}'] for i in range(1, 21)]) else '',
           f'=IFERROR(INDEX(Ringkasan!$U:$U,MATCH(C{i},Ringkasan!$D:$D,0)),"Tidak")'] for i, r in enumerate(kc, 2)]
wsK = sheet('Kecemasan', k_kol, k_rows, {1: 16, 2: 10, 3: 14, 7: 21})
fmt(wsK, 'G', 'yyyy-mm-dd hh:mm', 2, n)
wsK.conditional_formatting.add(f'M2:AF{n}', ColorScaleRule(start_type='num', start_value=1, start_color='63BE7B', mid_type='num', mid_value=3, mid_color='FFEB84', end_type='num', end_value=5, end_color='F8696B'))
wsK.conditional_formatting.add(f'H2:H{n}', ColorScaleRule(start_type='num', start_value=20, start_color='63BE7B', mid_type='num', mid_value=60, mid_color='FFEB84', end_type='num', end_value=100, end_color='F8696B'))
wsK.column_dimensions['A'].hidden = True
for j in range(13, 33): wsK.column_dimensions[L(j)].width = 5

# ---------- Tes ----------
ts = sorted(tes, key=lambda r: (r['uid'], r['fase'] != 'pre'))
t_kol = ['Kunci', 'Nama', 'uid', 'Putaran', 'Level', 'Fase', ('Paket', 'Paket A/B diacak per peserta (AB atau BA) supaya tes awal ≠ tes akhir.'),
         ('Skor %', 'Gabungan 聽力 + 閱讀'), ('聽力 %', 'Soal q1–q10'), ('閱讀 %', 'Soal q11–q20'), 'Benar', 'Dari', ('Menit', 'Lama mengerjakan (batas 30 menit)'),
         'Selesai', ('Versi tes', 'Lama = sebelum 8 Okt 2026 00:19 (kunci kebanyakan A, pilihan belum diacak). Baru = pilihan diacak.')] + \
        [(f'q{i}', '1 = benar, 0 = salah') for i in range(1, 21)] + \
        [(f'j{i}', 'Huruf yang dipilih, menurut urutan ASLI di data soal') for i in range(1, 21)] + ['Pakai?']
n = len(ts) + 1
t_rows = [[f"{r['uid']}|{r['fase']}", r['nama'], r['uid'], angka(r['putaran']), r['level'], 'awal' if r['fase'] == 'pre' else 'akhir', r['paket'],
           angka(r['skor']), angka(r['dengar']), angka(r['baca']), angka(r['benar']), angka(r['dari']),
           round((angka(r['detik']) or 0) / 60, 1), tgl(r['selesai']), 'Lama' if (tgl(r['selesai']) or TES_BARU) < TES_BARU else 'Baru'] +
          [angka(r[f'q{i}']) for i in range(1, 21)] + [r[f'j{i}'] for i in range(1, 21)] +
          [f'=IFERROR(INDEX(Ringkasan!$U:$U,MATCH(C{i},Ringkasan!$D:$D,0)),"Tidak")'] for i, r in enumerate(ts, 2)]
wsT = sheet('Tes', t_kol, t_rows, {1: 16, 3: 14, 14: 21})
fmt(wsT, 'N', 'yyyy-mm-dd hh:mm', 2, n)
wsT.conditional_formatting.add(f'P2:AI{n}', CellIsRule(operator='equal', formula=['1'], fill=HIJAU))
wsT.conditional_formatting.add(f'P2:AI{n}', CellIsRule(operator='equal', formula=['0'], fill=MERAH))
wsT.conditional_formatting.add(f'O2:O{n}', CellIsRule(operator='equal', formula=['"Lama"'], fill=MERAH))
wsT.column_dimensions['A'].hidden = True
for j in range(16, 56): wsT.column_dimensions[L(j)].width = 4.5

# ---------- Modul ----------
md = sorted(modul, key=lambda r: (r['uid'], r['kode']))
m_kol = ['Nama', 'uid', 'Putaran', 'Level', 'Bab', ('Tahap', 'Tahap yang sedang dikerjakan (1–6)'), 'Selesai', 'Waktu selesai',
         ('Tugas terbaik %', 'Nilai 溝通任務 terbaik'), 'Tugas dicoba', ('Kata hafal', 'Jumlah kata yang ditandai "Sudah hafal"'),
         ('Target yakin', 'Jumlah can-do yang dinilai 😎 Yakin'), ('Menit belajar', 'Hanya saat layar aktif & ada sentuhan dalam 3 menit'),
         'Target pribadi', 'Bagian yang masih sulit']
m_rows = [[r['nama'], r['uid'], angka(r['putaran']), r['level'], r['kode'], angka(r['tahap']), 'Ya' if r['selesai'] == 'ya' else '',
           tgl(r['selesai_waktu']), angka(r['tugas_best']), angka(r['tugas_kali']), angka(r['kata_hafal']), angka(r['yakin']),
           angka(r['menit_belajar']), r['target'], r['sulit']] for r in md]
wsM = sheet('Modul', m_kol, m_rows, {2: 14, 8: 21, 14: 30, 15: 30})
fmt(wsM, 'H', 'yyyy-mm-dd hh:mm', 2, len(md) + 1)

# ---------- Evaluasi ----------
ev = sorted(evals, key=lambda r: r['uid'])
e_kol = ['Nama', 'uid', 'Putaran', 'Level', 'Waktu'] + [(f'e{i}', t) for i, t in enumerate(EVAL, 1)] + \
        [('Rata-rata', '1 = sangat tidak setuju … 5 = sangat setuju'), 'Seragam?'] + [(f'Terbuka {i}', t) for i, t in enumerate(KUES['terbuka'], 1)] + ['Pakai?']
n = len(ev) + 1
e_rows = [[r['nama'], r['uid'], angka(r['putaran']), r['level'], tgl(r['waktu'])] + [angka(r[f'e{k}']) for k in range(1, 11)] +
          [f'=AVERAGE(F{i}:O{i})', 'Ya' if seragam([r[f'e{k}'] for k in range(1, 11)]) else ''] + [r[f'terbuka{k}'] for k in range(1, 4)] +
          [f'=IFERROR(INDEX(Ringkasan!$U:$U,MATCH(B{i},Ringkasan!$D:$D,0)),"Tidak")'] for i, r in enumerate(ev, 2)]
wsE = sheet('Evaluasi', e_kol, e_rows, {2: 14, 5: 21, 18: 35, 19: 35, 20: 35})
fmt(wsE, 'E', 'yyyy-mm-dd hh:mm', 2, n); fmt(wsE, 'P', '0.0', 2, n)
for j in range(6, 16): wsE.column_dimensions[L(j)].width = 5
for r in range(2, n + 1):
    for c in 'RST': wsE[f'{c}{r}'].alignment = Alignment(wrap_text=True, vertical='top')

# ---------- Ringkasan ----------
R0 = 12   # baris judul tabel
ps = sorted(peserta, key=lambda p: (tgl(p['mulai']) or datetime.max, int(angka(p['putaran']) or 1)))
r_kol = ['No', 'Nama', 'Email', 'uid', 'Putaran', 'Level', 'Perangkat', 'Bahasa', 'Mulai', 'Status',
         ('Cemas awal', '20–100, makin tinggi makin cemas'), 'Cemas akhir', ('Δ Cemas', 'Akhir − awal. NEGATIF = kecemasan turun (baik)'),
         'Tes awal %', 'Tes akhir %', ('Δ Tes', 'Akhir − awal. POSITIF = nilai naik (baik)'), 'Bab selesai', 'Menit belajar', ('Evaluasi', 'Rata-rata e1–e10 (1–5)'),
         ('Catatan otomatis', 'Alasan baris ini disarankan TIDAK dipakai. Dibuat oleh rekap_hasil.py.'),
         ('Pakai?', 'Ya/Tidak — saran otomatis, boleh diubah. Semua rata-rata di atas & di tab Per butir hanya menghitung baris "Ya".')]
r_rows = []
for i, p in enumerate(ps, R0 + 1):
    cat, pakai = catatan(p)
    look = lambda tab_, kol, fase: f'IFERROR(INDEX({tab_}!${kol}:${kol},MATCH(D{i}&"|{fase}",{tab_}!$A:$A,0)),"")'
    r_rows.append([i - R0, p['nama'], p['email'], p['uid'], angka(p['putaran']), p['level'], p['perangkat'], p['bahasa'], tgl(p['mulai']), status(p),
                   '=' + look('Kecemasan', 'H', 'pre'), '=' + look('Kecemasan', 'H', 'post'), f'=IF(AND(ISNUMBER(K{i}),ISNUMBER(L{i})),L{i}-K{i},"")',
                   '=' + look('Tes', 'H', 'pre'), '=' + look('Tes', 'H', 'post'), f'=IF(AND(ISNUMBER(N{i}),ISNUMBER(O{i})),O{i}-N{i},"")',
                   angka(p['bab_selesai']), f'=SUMIFS(Modul!$M:$M,Modul!$B:$B,D{i})',
                   f'=IFERROR(INDEX(Evaluasi!$P:$P,MATCH(D{i},Evaluasi!$B:$B,0)),"")', cat, pakai])
ws = sheet('Ringkasan', r_kol, r_rows, {2: 10, 3: 26, 4: 14, 9: 21, 10: 20, 20: 48}, mulai=R0)
wb.move_sheet('Ringkasan', -(len(wb.sheetnames) - 1))
last = R0 + max(1, len(r_rows))
fmt(ws, 'I', 'yyyy-mm-dd hh:mm', R0 + 1, last); fmt(ws, 'S', '0.0', R0 + 1, last)
for col in 'MP': fmt(ws, col, '+0;-0;0', R0 + 1, last)
for r in range(R0 + 1, last + 1):
    ws[f'U{r}'].fill = INPUT; ws[f'U{r}'].alignment = Alignment(horizontal='center')
    ws[f'T{r}'].alignment = Alignment(wrap_text=True, vertical='top')
dv = DataValidation(type='list', formula1='"Ya,Tidak"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'U{R0 + 1}:U{last}')
ws.conditional_formatting.add(f'M{R0 + 1}:M{last}', CellIsRule(operator='lessThan', formula=['0'], fill=HIJAU))
ws.conditional_formatting.add(f'M{R0 + 1}:M{last}', CellIsRule(operator='greaterThan', formula=['0'], fill=MERAH))
ws.conditional_formatting.add(f'P{R0 + 1}:P{last}', CellIsRule(operator='greaterThan', formula=['0'], fill=HIJAU))
ws.conditional_formatting.add(f'P{R0 + 1}:P{last}', CellIsRule(operator='lessThan', formula=['0'], fill=MERAH))

# blok ringkasan di atas tabel
ws['A1'] = 'Rekap uji coba modul mini TOCFL'; ws['A1'].font = Font(name=F, bold=True, size=14)
ws['A2'] = f'Sumber: {SRC.name} · dibuat {datetime.now():%Y-%m-%d %H:%M} · kolom kuning "Pakai?" boleh diubah (Ya/Tidak)'
ws['A2'].font = Font(name=F, italic=True, color='595959')
rng = lambda c: f'{c}${R0 + 1}:{c}${last}'
U = rng('$U')
stat = [('Jumlah baris (peserta × putaran)', f'=COUNTA({rng("$D")})', '0'),
        ('Dipakai untuk analisis (Pakai = Ya)', f'=COUNTIF({U},"Ya")', '0'),
        ('Rata-rata cemas awal → akhir', f'=IFERROR(AVERAGEIFS({rng("$K")},{U},"Ya"),"–")', f'=IFERROR(AVERAGEIFS({rng("$L")},{U},"Ya"),"–")'),
        ('Rata-rata tes awal → akhir (%)', f'=IFERROR(AVERAGEIFS({rng("$N")},{U},"Ya"),"–")', f'=IFERROR(AVERAGEIFS({rng("$O")},{U},"Ya"),"–")'),
        ('Rata-rata evaluasi modul (1–5)', f'=IFERROR(AVERAGEIFS({rng("$S")},{U},"Ya"),"–")', None)]
ws['D4'], ws['E4'] = 'Awal / nilai', 'Akhir'
for c in ('D4', 'E4'): ws[c].font = Font(name=F, bold=True, color='595959')
for k, (lbl, a, b) in enumerate(stat, 5):
    ws[f'A{k}'] = lbl; ws[f'A{k}'].font = Font(name=F)
    ws[f'D{k}'] = a; ws[f'D{k}'].font = Font(name=F, bold=True); ws[f'D{k}'].number_format = '0.0'
    if b and b != '0': ws[f'E{k}'] = b; ws[f'E{k}'].font = Font(name=F, bold=True); ws[f'E{k}'].number_format = '0.0'
    if b == '0': ws[f'D{k}'].number_format = '0'
ws['A10'] = 'Catatan: baris dengan "tes versi lama" dikerjakan sebelum perbaikan 8 Okt 2026 (kunci jawaban kebanyakan A) — nilai tesnya tidak sebanding dengan data baru.'
ws['A10'].font = Font(name=F, italic=True, color='C00000')

# ---------- Per butir ----------
wsB = wb.create_sheet('Per butir', 1)
wsB['A1'] = 'Rata-rata per butir (hanya baris Pakai = Ya)'; wsB['A1'].font = Font(name=F, bold=True, size=13)
kepala = lambda r, cols: [setattr(wsB.cell(r, j, v), 'fill', HEAD) or setattr(wsB.cell(r, j), 'font', HFONT) for j, v in enumerate(cols, 1)]
kepala(3, ['Butir', 'Dimensi', 'Pernyataan (kuesioner kecemasan)', 'Awal', 'Akhir', 'Δ', 'n awal', 'n akhir'])
nk = len(kc) + 1
for i, (d, t, r) in enumerate(BUTIR, 1):
    row, col = 3 + i, L(12 + i)   # b1 = kolom M di tab Kecemasan
    vals = [f'b{i}' + (' (R)' if r else ''), d.split('.')[0], t + (' — butir terbalik: makin tinggi = makin yakin' if r else ''),
            f'=IFERROR(AVERAGEIFS(Kecemasan!${col}$2:${col}${nk},Kecemasan!$F$2:$F${nk},"awal",Kecemasan!$AH$2:$AH${nk},"Ya"),"–")',
            f'=IFERROR(AVERAGEIFS(Kecemasan!${col}$2:${col}${nk},Kecemasan!$F$2:$F${nk},"akhir",Kecemasan!$AH$2:$AH${nk},"Ya"),"–")',
            f'=IF(AND(ISNUMBER(D{row}),ISNUMBER(E{row})),E{row}-D{row},"")',
            f'=COUNTIFS(Kecemasan!$F$2:$F${nk},"awal",Kecemasan!$AH$2:$AH${nk},"Ya")',
            f'=COUNTIFS(Kecemasan!$F$2:$F${nk},"akhir",Kecemasan!$AH$2:$AH${nk},"Ya")']
    for j, v in enumerate(vals, 1): wsB.cell(row, j, v).font = Font(name=F)
r2 = 3 + len(BUTIR) + 2
kepala(r2, ['Butir', '', 'Pernyataan (evaluasi modul)', 'Rata-rata', '', '', 'n'])
ne = len(ev) + 1
for i, t in enumerate(EVAL, 1):
    row, col = r2 + i, L(5 + i)
    vals = [f'e{i}', '', t, f'=IFERROR(AVERAGEIFS(Evaluasi!${col}$2:${col}${ne},Evaluasi!$U$2:$U${ne},"Ya"),"–")', '', '',
            f'=COUNTIFS(Evaluasi!$U$2:$U${ne},"Ya")']
    for j, v in enumerate(vals, 1): wsB.cell(row, j, v).font = Font(name=F)
for r in range(4, r2 + len(EVAL) + 1):
    for c in 'DEF': wsB[f'{c}{r}'].number_format = '0.00'
    wsB[f'C{r}'].alignment = Alignment(wrap_text=True, vertical='top')
wsB.conditional_formatting.add(f'F4:F{3 + len(BUTIR)}', CellIsRule(operator='lessThan', formula=['0'], fill=HIJAU))
wsB.conditional_formatting.add(f'F4:F{3 + len(BUTIR)}', CellIsRule(operator='greaterThan', formula=['0'], fill=MERAH))
for c, w in zip('ABCDEFGH', (8, 9, 70, 9, 9, 9, 8, 8)): wsB.column_dimensions[c].width = w
wsB.freeze_panes = 'A4'

# ---------- Keterangan ----------
wsI = wb.create_sheet('Keterangan')
teks = [('Cara memperbarui', True),
        ('1. Google Sheet → File → Download → Microsoft Excel (.xlsx), simpan sebagai tocfl-mini/測試結果.xlsx (timpa yang lama).', False),
        ('2. Jalankan:  py TOCFL/_kerja/mini/rekap_hasil.py   → file ini dibuat ulang. Ubahan manual di kolom "Pakai?" akan hilang.', False),
        ('3. File ini berisi email peserta: jangan di-commit / dibagikan ke publik.', False), ('', False),
        ('Tab', True),
        ('Ringkasan — satu baris per peserta per putaran: cemas & tes awal/akhir, perubahan (Δ), status, dan kolom "Pakai?".', False),
        ('Per butir — rata-rata tiap butir kuesioner kecemasan (awal vs akhir) dan tiap pernyataan evaluasi; hanya baris Pakai = Ya.', False),
        ('Kecemasan / Tes / Modul / Evaluasi — data mentah yang dirapikan. Arahkan kursor ke judul kolom untuk melihat pernyataan lengkap.', False), ('', False),
        ('Cara membaca', True),
        ('Kecemasan: total 20–100, makin tinggi makin cemas. Δ negatif (hijau) = kecemasan turun. Dimensi A–D masing-masing 5–25.', False),
        ('b10 adalah butir terbalik ("Saya yakin…"): di kolom b10 nilainya mentah; di Total dan dimensi B sudah dibalik (6 − nilai).', False),
        ('Tes: skor dalam %, 聽力 = q1–q10, 閱讀 = q11–q20. Δ positif (hijau) = nilai naik. q = 1 benar / 0 salah; j = huruf yang dipilih (urutan asli).', False),
        ('Evaluasi: 1 = sangat tidak setuju … 5 = sangat setuju.', False), ('', False),
        ('Aturan otomatis "Pakai? = Tidak"', True),
        ('• putaran 2 (analisis utama memakai putaran 1)   • belum selesai (belum tes akhir + evaluasi)', False),
        ('• email sama dengan baris lain dari perangkat lain (yang dipakai: perangkat yang mulai paling awal)', False),
        ('• jawaban seragam (semua butir kuesioner/evaluasi diberi nilai yang sama — kemungkinan asal isi)', False),
        (f'• tes versi lama: dikerjakan sebelum {TES_BARU:%d %b %Y %H:%M}, ketika kunci jawaban kebanyakan A dan pilihan belum diacak', False)]
for i, (t, b) in enumerate(teks, 1):
    wsI.cell(i, 1, t).font = Font(name=F, bold=b, size=12 if b else 10)
wsI.column_dimensions['A'].width = 130

del wb['Sheet']
wb.active = 0
wb.save(OUT)
print('OK', OUT, '·', len(ps), 'baris peserta')
