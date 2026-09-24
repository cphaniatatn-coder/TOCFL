import json, sys, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
S=sys.argv[1]; LV=[int(x) for x in sys.argv[2].split(',')]
tb=json.load(open(S+'/tbcl.json',encoding='utf-8')); G=json.load(open(S+'/grammar.json',encoding='utf-8'))
py={}
for o in tb:
    for p in [o['w']]+o['w'].split('/'): py.setdefault((o['lv'],p),o['py'])
CAT={1:'個人資料',2:'日常起居',3:'職業',4:'休閒、娛樂',5:'交通、旅遊',6:'社交、人際',7:'身體、醫療',8:'教育、學習',9:'購物、商店',10:'餐飲、烹飪',11:'公共服務',12:'安全',13:'自然環境',14:'社會',15:'文化',16:'情緒、態度',17:'科技'}
VOL={1:'Vol.1 · A0 (TBCL 第1級)',2:'Vol.2 · A1 (TBCL 第2級)',3:'Vol.3 · A2 (TBCL 第3級)'}
wb=openpyxl.Workbook(); ws0=wb.active; ws0.title='Ringkasan & rumus'
H=PatternFill('solid',fgColor='1F4E78')
def hdr(ws):
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=H; c.alignment=Alignment(wrap_text=True,vertical='center')
ws0.append(['Rumus jumlah modul (turunan CLT + cakupan penuh)']); ws0['A1'].font=Font(bold=True,size=13)
ws0.append(['N = maks( ⌈jumlah poin grammar ÷ 2⌉ , ⌈jumlah kata level ÷ 18⌉ )   — 2 = batas grammar Sweller; 18 = 10 核心 + ±8 補充'])
ws0.append([]); ws0.append(['Level','Kata resmi','Grammar resmi','⌈G÷2⌉','⌈W÷18⌉','N minimum','N hasil desain','Catatan'])
import math
for lv,(w,g) in {1:(396,15),2:(402,92),3:(456,134)}.items():
    n=None
    try: n=sum(1 for l in open(f'{S}/plan_L{lv}.txt',encoding='utf-8') if l.strip() and not l.startswith('#'))
    except FileNotFoundError: pass
    note={1:'A0 vocabulary-first: kosakata yang menentukan. 171 工具詞 menaikkan kebutuhan slot → 25 modul (grammar diminimalkan, 10 modul tanpa grammar baru = daur ulang).',2:'Grammar yang menentukan → 新詞量 ±9/modul',3:'Grammar yang menentukan → 新詞量 ±7/modul'}[lv]
    ws0.append([VOL[lv],w,g,math.ceil(g/2),math.ceil(w/18),max(math.ceil(g/2),math.ceil(w/18)),n or '(belum dirancang)',note])
for c in ws0[4]: c.font=Font(bold=True)
ws0.column_dimensions['A'].width=26; ws0.column_dimensions['H'].width=90
for lv in LV:
    ws=wb.create_sheet(VOL[lv].split(' (')[0].replace('·','-'))
    ws.append(['ID','Judul (ujaran nyata)','Adegan','Kategori 情境','Grammar baru','核心 (produktif)','補充 (reseptif)','#核心','#補充','新詞量'])
    hdr(ws)
    gm={int(x['id']):x['g'] for x in G[f'第{lv}級']}
    for l in open(f'{S}/plan_L{lv}.txt',encoding='utf-8'):
        if not l.strip() or l.startswith('#'): continue
        f=l.rstrip('\n').split('|')
        core=f[5].split(); sup=f[6].split() if len(f)>6 else []
        fmt=lambda ws_: '、'.join(f"{w.split('@')[0]}（{py.get((lv,w.split('@')[0]),'?')}）"+(f"［義項: {w.split('@')[1]}］" if '@' in w else '') for w in ws_)
        cats=' + '.join(f'{c}.{CAT[int(c)]}' for c in f[3].split(','))
        gram='—（daur ulang）' if f[4] in ('—','') else '；'.join(f'{x}. {gm[int(x)]}' for x in f[4].split(','))
        ws.append([f[0],f[1],f[2],cats,gram,fmt(core),fmt(sup),len(core),len(sup),len(core)+len(sup)])
    for col,wd in zip('ABCDEFGHIJ',[6,24,34,20,22,60,55,7,7,8]): ws.column_dimensions[col].width=wd
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.freeze_panes='C2'
wb.save(r'C:\Users\user\OneDrive\文件\TOCFL\rencana-modul.xlsx'); print('ok')
